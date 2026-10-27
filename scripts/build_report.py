"""Build the final monochrome double-spaced academic report from Markdown."""
from pathlib import Path
import json
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'deliverables'
BLACK=RGBColor(0,0,0)

def inline(p,text):
    for part in re.split(r'(\*\*.*?\*\*|\*[^*]+\*|`[^`]+`)',text):
        r=p.add_run(part.strip('*`') if part.startswith(('*','`')) else part)
        r.font.name='Times New Roman';r.font.size=Pt(12);r.font.color.rgb=BLACK
        if part.startswith('**'):r.bold=True
        elif part.startswith('*'):r.italic=True

def no_border(element):
    for border in list(element.findall('.//'+qn('w:pBdr'))):border.getparent().remove(border)

def body_paragraph(doc,text,reference=False):
    p=doc.add_paragraph();inline(p,text)
    p.paragraph_format.first_line_indent=Inches(-.5 if reference else .5)
    if reference:p.paragraph_format.left_indent=Inches(.5)
    return p

def table(doc,rows):
    t=doc.add_table(rows=0,cols=len(rows[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    widths=([1.55]+[(6.27-1.55)/(len(rows[0])-1)]*(len(rows[0])-1))
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        border=OxmlElement('w:'+side);border.set(qn('w:val'),'single');border.set(qn('w:sz'),'4');border.set(qn('w:color'),'D9D9D9');borders.append(border)
    t._tbl.tblPr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for name in ['top','left','bottom','right']:
        el=OxmlElement('w:'+name);el.set(qn('w:w'),'85');el.set(qn('w:type'),'dxa');margins.append(el)
    t._tbl.tblPr.append(margins)
    for j,row in enumerate(rows):
        cells=t.add_row().cells
        pr=t.rows[j]._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if j==0:pr.append(OxmlElement('w:tblHeader'))
        for k,(c,txt) in enumerate(zip(cells,row)):
            c.width=Inches(widths[k]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=c.paragraphs[0];p.paragraph_format.first_line_indent=Inches(0)
            p.paragraph_format.line_spacing=2;p.paragraph_format.space_after=Pt(0)
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT if k==0 else WD_ALIGN_PARAGRAPH.CENTER
            inline(p,txt)
            if j==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EEEEEE');c._tc.get_or_add_tcPr().append(shade)
                for r in p.runs:r.bold=True
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

def count_words(text):
    text=re.sub(r'!\[([^\]]*)\]\([^)]*\)',r'\1',text)
    text=re.sub(r'^\|[\s:|\-]+\|\s*$','',text,flags=re.M)
    return len(re.findall(r"\b[\w]+(?:[-'][\w]+)*\b",text))

def build():
    OUT.mkdir(exist_ok=True);(ROOT/'build').mkdir(exist_ok=True)
    source=(ROOT/'docs/report.md').read_text(encoding='utf-8')
    doc=Document();sec=doc.sections[0]
    sec.page_height=Inches(11.69);sec.page_width=Inches(8.27)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
    sec.header_distance=sec.footer_distance=Inches(.5)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
        st=doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(12);st.font.color.rgb=BLACK
        fonts=st.element.get_or_add_rPr().get_or_add_rFonts()
        for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
            fonts.attrib.pop(qn('w:'+key),None)
        for key in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+key),'Times New Roman')
        st.paragraph_format.line_spacing=2;st.paragraph_format.space_after=Pt(0)
        st.paragraph_format.widow_control=True;no_border(st.element)
    for name in ['Heading 1','Heading 2','Heading 3']:
        st=doc.styles[name];st.font.bold=True;st.paragraph_format.space_before=Pt(12);st.paragraph_format.keep_with_next=True
    doc.styles['Title'].font.size=Pt(16);doc.styles['Title'].font.bold=True
    doc.styles['Caption'].font.italic=False
    doc.core_properties.title='CampusCompass Evidence Grounded University Policy Assistance'
    doc.core_properties.author='Ayman Lakhnati'
    p=doc.add_paragraph('CampusCompass Evidence Grounded University Policy Assistance','Title');p.alignment=1;p.paragraph_format.keep_with_next=True
    for text in ['Ayman Lakhnati','MAAI 902 Natural Language Processing and LLMs','Assessment Task 2 NLP Solution Design','26 October 2026']:
        p=doc.add_paragraph(text);p.alignment=1;p.paragraph_format.keep_with_next=True
    main=re.search(r'^## Abstract\s*$',source,re.M)
    lines=source[main.start():].splitlines();i=0;references=False
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line:continue
        if line.startswith('## '):
            title=re.sub(r'[^\w\s]','',line[3:])
            if title.lower()=='references':doc.add_page_break();references=True
            p=doc.add_paragraph(title,'Heading 1')
            continue
        if line.startswith('### '):doc.add_paragraph(re.sub(r'[^\w\s]','',line[4:]),'Heading 2');continue
        if re.match(r'^Table \d+ ',line):
            p=doc.add_paragraph(line,'Caption');p.paragraph_format.keep_with_next=True
            continue
        if line.startswith('|') and i<len(lines) and re.match(r'^\|[\s:|\-]+\|$',lines[i].strip()):
            rows=[[s.strip() for s in line.strip('|').split('|')]];i+=1
            while i<len(lines) and lines[i].strip().startswith('|'):
                rows.append([s.strip() for s in lines[i].strip().strip('|').split('|')]);i+=1
            table(doc,rows);continue
        if line.startswith('!['):
            match=re.match(r'!\[(.*?)\]\((.*?)\)',line)
            if match:
                path=ROOT/match.group(2)
                if not path.exists():raise FileNotFoundError(path)
                p=doc.add_paragraph();p.alignment=1;p.paragraph_format.keep_with_next=True;p.paragraph_format.line_spacing=1
                p.add_run().add_picture(str(path),width=Inches(6.2))
                p=doc.add_paragraph(match.group(1),'Caption');p.paragraph_format.keep_with_next=False
            continue
        p=body_paragraph(doc,line,reference=references)
        if not references and len(doc.paragraphs)>1:
            prev=doc.paragraphs[-2]
            if prev.style.name.startswith('Heading'):p.paragraph_format.first_line_indent=Inches(0)
    footer=sec.footer.paragraphs[0];footer.alignment=1
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    footer.style=doc.styles['Normal']
    for p in doc.paragraphs:
        no_border(p._p)
        for r in p.runs:
            r.font.name='Times New Roman';r.font.color.rgb=BLACK
            r.font.size=Pt(16 if p.style.name=='Title' else 12)
            fonts=r._element.get_or_add_rPr().get_or_add_rFonts()
            for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:fonts.attrib.pop(qn('w:'+key),None)
            for key in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+key),'Times New Roman')
    path=OUT/'CampusCompass_Final_Report.docx';doc.save(path)
    before_refs=re.split(r'^## References\s*$',source,flags=re.M)[0]
    assessed=re.search(r'^## 1[ .]',before_refs,re.M)
    main_body=before_refs[assessed.start():]
    main_body=re.split(r'^## AI assistance disclosure\s*$',main_body,flags=re.M)[0]
    receipt={'main_body_words':count_words(main_body),'abstract_through_disclosure_words':count_words(before_refs[main.start():]),'source':'docs/report.md','output':path.name,'format':'A4 Times New Roman 12 pt double spaced 1 inch margins monochrome','count_method':'Alphanumeric words including section headings table contents and figure captions; excludes title metadata abstract disclosure references for main_body_words; Markdown syntax and figure paths excluded.'}
    if not 3000<=receipt['main_body_words']<=4000:raise ValueError(receipt)
    (ROOT/'build/final_report_receipt.json').write_text(json.dumps(receipt,indent=2),encoding='utf-8');print(json.dumps(receipt,indent=2))

if __name__=='__main__':build()
